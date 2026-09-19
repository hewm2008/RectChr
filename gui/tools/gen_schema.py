#!/usr/bin/env python3
"""Generate gui/resources/params_schema.json for the RectChr GUI.

Sources:
  1. NewParaList.xlsx  (sheet 参数列表 = main metadata, sheet cheeckAA = zh descriptions)
  2. Code cross-validation: keys actually read from %HashConfi in bin/lib/*.pm + bin/RectChr
  3. Coded defaults parsed from ConfigDataLoader::init_global_defaults
Run from anywhere:  python3 gui/tools/gen_schema.py [repo_root]
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2]
XLSX = ROOT / "NewParaList.xlsx"
OUT = ROOT / "gui" / "resources" / "params_schema.json"

# ---------------------------------------------------------------- categories
CATEGORIES = {
    "files":        {"zh": "数据文件",       "en": "Data files"},
    "canvas":       {"zh": "画布",           "en": "Canvas"},
    "chr":          {"zh": "染色体",         "en": "Chromosomes"},
    "title":        {"zh": "标题",           "en": "Title"},
    "axis":         {"zh": "坐标轴",         "en": "Axes"},
    "track":        {"zh": "轨道通用",       "en": "Track general"},
    "data":         {"zh": "数据处理",       "en": "Data processing"},
    "label":        {"zh": "轨道名称",       "en": "Track label"},
    "geometry":     {"zh": "几何属性",       "en": "Geometry"},
    "color_legend": {"zh": "颜色与图例",     "en": "Colors & legend"},
    "svg":          {"zh": "样式 (SVG)",     "en": "Style (SVG)"},
}
CAT_MAP = {
    "画板": "canvas", "染色体": "chr", "title": "title", "刻度尺": "axis",
    "track": "track", "track-数据处理": "data", "track-Name_ID": "label",
    "track-sub-属性": "geometry", "stack-几何映射": "geometry",
    "track-颜色属性-图例": "color_legend", "track-sub-sVG属性": "svg",
}

# ---------------------------------------------------------------- plot types
PLOT_TYPES = [
    {"name": "heatmap",           "aliases": [],                "zh": "热图",       "en": "Heatmap"},
    {"name": "highlights",        "aliases": [],                "zh": "高亮",       "en": "Highlights"},
    {"name": "point",             "aliases": ["scatter", "points"], "zh": "点图",   "en": "Scatter / point"},
    {"name": "line",              "aliases": ["lines"],         "zh": "线图",       "en": "Line"},
    {"name": "hist",              "aliases": ["histogram"],     "zh": "柱状图",     "en": "Histogram"},
    {"name": "text",              "aliases": [],                "zh": "文本",       "en": "Text"},
    {"name": "shape",             "aliases": ["shapes", "Shape"], "zh": "形状",     "en": "Shape"},
    {"name": "ridgeline",         "aliases": [],                "zh": "山脊线",     "en": "Ridgeline"},
    {"name": "PairWiseLink",      "aliases": ["pairwiselink"],  "zh": "彩虹链接",   "en": "Pairwise link"},
    {"name": "PairWiseLinkV2",    "aliases": ["pairwiselinkV2"], "zh": "彩虹链接V2", "en": "Pairwise link V2"},
    {"name": "LinkS",             "aliases": ["LinkSelf"],      "zh": "自连接",     "en": "Self link"},
    {"name": "heatmapAnimated",   "aliases": [],                "zh": "动态热图",   "en": "Animated heatmap"},
    {"name": "histAnimated",      "aliases": ["histogramAnimated"], "zh": "动态柱状图", "en": "Animated histogram"},
    {"name": "highlightsAnimated","aliases": [],                "zh": "动态高亮",   "en": "Animated highlights"},
    {"name": "link",              "aliases": ["Link"],          "zh": "Link(遗留)", "en": "Link (legacy)",
     "legacy": True},
]
VALUE_PLOTS = ["point", "line", "hist", "ridgeline"]
LINK_PLOTS = ["PairWiseLink", "PairWiseLinkV2", "LinkS", "link"]

# param -> plot types it primarily applies to (empty = all)
RELEVANCE = {
    "track_geom_shape": ["shape"], "track_geom_shape_size": ["shape"],
    "track_point_size": ["point"],
    "track_text_size": ["text"], "track_text_angle": ["text"],
    "track_text_anchor": ["text"], "track_text_overlap": ["text"],
    "link_direction": LINK_PLOTS, "link_linestyle": LINK_PLOTS,
    "link_uniform_height": LINK_PLOTS, "line_colors_conf": LINK_PLOTS,
    "cutoff_y": VALUE_PLOTS, "cutoff_color": VALUE_PLOTS,
    "cutoff1_y": VALUE_PLOTS, "cutoff1_color": VALUE_PLOTS,
    "cutoff2_y": VALUE_PLOTS, "cutoff2_color": VALUE_PLOTS,
    "log_p": VALUE_PLOTS,
    "yaxis_tick_show": VALUE_PLOTS + ["shape"],
    "yaxis_tick_num": VALUE_PLOTS + ["shape"], "yaxis_tick_precision": VALUE_PLOTS + ["shape"],
    "yaxis_tick_unit": VALUE_PLOTS + ["shape"], "yaxis_tick_strokewidth": VALUE_PLOTS + ["shape"],
    "yaxis_tick_direction": VALUE_PLOTS + ["shape"],
    "as_flag": ["point", "shape", "text", "heatmap", "highlights"],
}

# known enum choices
CHOICES = {
    "chr_orientation": ["vertical", "horizontal"],
    "plot_type": [p["name"] for p in PLOT_TYPES],
    "link_direction": ["UpDown", "DownUp", "UpUp", "DownDown", "UpDownV2", "DownUpV2"],
    "xaxis_tick_direction": ["Down", "Up"],
    "yaxis_tick_direction": ["Down", "Up"],
    "track_text_anchor": ["start", "middle", "end"],
    "colormap_legend_layout": [str(i) for i in range(0, 14)],
}

# English labels/descriptions (concise, hand written)
EN = {
    "File1": ("Data file 1", "Main input data file, format: Chr Start End Value1 ... (NA = skip)"),
    "FileX": ("Data file N", "Optional extra data file (File2, File3, ...)"),
    "track_num": ("Track count", "Number of tracks (layers); auto-inferred from File1 columns by default"),
    "show_columns": ("Show columns", "Columns to plot, e.g. File1:4 or File2:4,5"),
    "plot_type": ("Plot type", "Drawing style of this track"),
    "title": ("Title text", "Main figure title"),
    "title_color": ("Title color", "Color of the main title"),
    "title_size": ("Title size ratio", "Font size ratio of the main title"),
    "title_shift_x": ("Title shift X", "Horizontal offset of the title"),
    "title_shift_y": ("Title shift Y", "Vertical offset of the title"),
    "chr_orientation": ("Chr orientation", "vertical or horizontal chromosome layout"),
    "chr_order": ("Chr order", "Chromosome order / filter list, e.g. chr1,chr2"),
    "chr_order_reverse": ("Reverse chr order", "Reverse the chromosome display order"),
    "chr_zoom_region": ("Zoom region", "Zoom into a region, format chr2:1000:5000"),
    "chr_spacing_ratio": ("Chr spacing ratio", "Spacing ratio between chromosomes"),
    "chr_label_rotation": ("Chr label rotation", "Rotation angle of chromosome labels"),
    "chr_label_size_ratio": ("Chr label size ratio", "Font size ratio of chromosome labels"),
    "chr_label_shift_x": ("Chr label shift X", "Horizontal shift of chromosome labels"),
    "chr_label_shift_y": ("Chr label shift Y", "Vertical shift of chromosome labels"),
    "chr_scale_ratio": ("Chr scale ratio", "Scale chromosome lengths by this ratio"),
    "canvas_body": ("Canvas body", "Main canvas size (default 1200)"),
    "canvas_margin_top": ("Canvas margin top", "Top margin of the canvas (default 55)"),
    "canvas_margin_bottom": ("Canvas margin bottom", "Bottom margin of the canvas (default 30)"),
    "canvas_margin_left": ("Canvas margin left", "Left margin of the canvas (default 100)"),
    "canvas_margin_right": ("Canvas margin right", "Right margin of the canvas (default 80)"),
    "canvas_height_ratio": ("Canvas height ratio", "Overall canvas height ratio"),
    "canvas_width_ratio": ("Canvas width ratio", "Overall canvas width ratio"),
    "canvas_angle": ("Canvas angle", "Rotation angle of the exported PNG (degrees)"),
    "colormap_conf": ("Colormap file", "Custom value-to-color mapping file, e.g. P1=\"#FE0808\""),
    "colormap_brewer_name": ("Brewer palette", "Preset palette name, e.g. GnYlRd (numeric) or Paired (categorical)"),
    "colormap_reverse": ("Reverse palette", "Reverse the color gradient order"),
    "colormap_low_color": ("Gradient low color", "Color for the lowest value"),
    "colormap_mid_color": ("Gradient mid color", "Color for the middle value"),
    "colormap_high_color": ("Gradient high color", "Color for the highest value"),
    "colormap_nlevels": ("Color levels", "Number of color gradient levels"),
    "colormap_legend_show": ("Show color legend", "Show (1) or hide (0) the gradient legend"),
    "colormap_legend_sizeratio": ("Legend size ratio", "Size ratio of the gradient legend (0 hides it)"),
    "colormap_legend_shift_x": ("Legend shift X", "Horizontal shift of the gradient legend"),
    "colormap_legend_shift_y": ("Legend shift Y", "Vertical shift of the gradient legend"),
    "colormap_legend_layout": ("Legend layout", "Legend layout mode (0-13)"),
    "colormap_legend_gap": ("Legend gap", "Gap between legend columns (global)"),
    "colormap_gradient_gap": ("Gradient gap", "Gap between adjacent gradient blocks"),
    "track_height": ("Track height", "Height of this track"),
    "track_bg_height_ratio": ("Track bg height ratio", "Background height as a proportion of track height"),
    "track_shift_x": ("Track shift X", "Horizontal shift of the whole track"),
    "track_shift_y": ("Track shift Y", "Vertical shift of the whole track"),
    "padding_ratio": ("Track padding ratio", "Vertical spacing ratio between adjacent tracks"),
    "label": ("Track label", "Label text of this track"),
    "label_color": ("Label color", "Color of the track label"),
    "label_size": ("Label size", "Font size ratio of the track label"),
    "label_angle": ("Label angle", "Rotation angle of the track label"),
    "label_shift_x": ("Label shift X", "Horizontal shift of the track label"),
    "label_shift_y": ("Label shift Y", "Vertical shift of the track label"),
    "background_color": ("Background color", "Background color of this track"),
    "background_show": ("Show background", "Show (1) or hide (0) the track background"),
    "bg_end_arc": ("Bg end arc", "Draw arcs at chromosome ends of the background"),
    "bg_end_offset": ("Bg end offset", "End curve starts at X=0 or afterwards"),
    "bg_end_arc_division": ("Bg end arc division", "Division of the end curve arc"),
    "bg_stroke_color": ("Bg stroke color", "Stroke color of the background"),
    "bg_stroke_width": ("Bg stroke width", "Stroke width of the background"),
    "upper_outlier_ratio": ("Upper outlier ratio", "Values above this quantile use the max color"),
    "lower_outlier_ratio": ("Lower outlier ratio", "Values below this quantile use the min color"),
    "Ymax": ("Y max", "Manual maximum of the display range"),
    "Ymin": ("Y min", "Manual minimum of the display range"),
    "cap_max_value": ("Cap max value", "Truncate values above this maximum"),
    "cap_min_value": ("Cap min value", "Truncate values below this minimum"),
    "log_p": ("-log10 value", "Apply -log10 transformation to values"),
    "as_flag": ("Treat as flag", "Treat the numeric column as ASCII flag (categorical)"),
    "axis_tick_unit": ("Axis tick unit", "Unit of the coordinate axis ticks (bp/Kb/Mb/...)"),
    "axis_text_angle": ("Axis text angle", "Rotation angle of the axis tick text"),
    "axis_tick_num": ("Axis tick number", "Number of axis ticks (max 10 parts by default)"),
    "axis_tick_interval": ("Axis tick interval", "Interval between axis ticks"),
    "axis_tick_precision": ("Axis tick precision", "Decimal digits of axis tick labels"),
    "xaxis_tick_show": ("Show X ticks", "Show (1) or hide (0) X axis tick labels"),
    "xaxis_tick_direction": ("X tick direction", "Draw X ticks Up or Down"),
    "xaxis_tick_color": ("X tick color", "Color of the X axis ticks"),
    "xaxis_shift_y": ("X axis shift Y", "Vertical shift of the X axis"),
    "yaxis_tick_show": ("Show Y ticks", "Show (1) or hide (0) Y axis tick labels"),
    "yaxis_tick_num": ("Y tick number", "Number of Y axis ticks"),
    "yaxis_tick_precision": ("Y tick precision", "Decimal digits of Y tick labels"),
    "yaxis_tick_unit": ("Y tick unit", "Unit of the Y axis ticks"),
    "yaxis_tick_strokewidth": ("Y tick stroke width", "Stroke width of the Y axis"),
    "yaxis_tick_direction": ("Y tick direction", "Draw Y ticks Up or Down"),
    "track_geom_shape": ("Shape type", "Shape type for shape plots (0-13)"),
    "track_geom_shape_size": ("Shape size ratio", "Size ratio of shapes"),
    "track_point_size": ("Point size ratio", "Size ratio of points"),
    "track_text_size": ("Text size ratio", "Font size ratio of texts"),
    "track_text_angle": ("Text angle", "Rotation angle of texts"),
    "track_text_anchor": ("Text anchor", "Text anchor (start/middle/end)"),
    "track_text_overlap": ("Text overlap", "0 = allow overlap; N = at most N overlapping texts"),
    "line_colors_conf": ("Line colors file", "Custom line color configuration file"),
    "link_direction": ("Link direction", "Direction of link start/end (UpDown, DownUp, ...)"),
    "link_linestyle": ("Link line style", "Straight line or bezier curve"),
    "link_uniform_height": ("Link uniform height", "Give all links the same height"),
    "cutoff_y": ("Cutoff line Y", "Y position of the cutoff line (docs name: cutoff_value)"),
    "cutoff_color": ("Cutoff line color", "Color of the cutoff line"),
    "cutoff1_y": ("Cutoff line 1 Y", "Y position of the first cutoff line"),
    "cutoff1_color": ("Cutoff line 1 color", "Color of the first cutoff line"),
    "cutoff2_y": ("Cutoff line 2 Y", "Y position of the second cutoff line"),
    "cutoff2_color": ("Cutoff line 2 color", "Color of the second cutoff line"),
    "stroke-width": ("Stroke width", "SVG stroke-width"),
    "strokewidth": ("Stroke width", "SVG stroke-width (alias)"),
    "stroke-opacity": ("Stroke opacity", "SVG stroke-opacity"),
    "fill": ("Fill color", "SVG fill color"),
    "fill-opacity": ("Fill opacity", "SVG fill-opacity"),
    "font-family": ("Font family", "SVG font-family"),
    "font-size": ("Font size", "SVG font-size"),
    "text-font-size": ("Text font size", "SVG text font size"),
    "hist_No_UD": ("Hist no U/D marks", "Internal: hide up/down marks for 2-column hist"),
}

# internal keys that must NOT be exposed as user parameters
INTERNAL = {
    "global", "ALL", "Bin", "RealBin", "IsNumber", "TotalValue", "MaxGradien",
    "Legend_OffsetRatio", "legend_Count", "legend_HHCount", "legend_HRCount",
    "legend_HVCount", "legend_VVCount", "legend_VRCount", "canvas_bodyOO",
    "canvas_margin_Height", "canvas_margin_Width", "ShiftChrNameRatio",
    "all_chromosomes_spacing", "NA", "zoom_region",
}

# xlsx rows that are NOT live parameters: the engine never reads these keys
# (verified) — they are legacy spellings, so they become aliases instead of
# user-facing params (see the alias table in main()).
DROP = {
    "strokewidth",            # legacy spelling of "stroke-width"
    "colormap_legend_size",   # legacy spelling of "colormap_legend_sizeratio"
}


def extract_palettes():
    """Palette catalog for the GUI picker.

    builtin: RColorBrewer names from ColorPaletteManager.pm (%MAX_COLOR_COUNT,
             %QualColNum) with exact preview colors parsed from %HashColData
             (max-level RGB triplets).
    files:   hex-per-line palettes from the ColorsBrewer/ directory."""
    import binascii

    pm = (ROOT / "bin" / "lib" / "ColorPaletteManager.pm").read_text(
        encoding="utf-8", errors="replace")
    maxc = {m.group(1): int(m.group(2)) for m in
            re.finditer(r'\$MAX_COLOR_COUNT\{"([^"]+)"\}\s*=\s*(\d+)', pm)}
    qual = {m.group(1) for m in
            re.finditer(r'\$QualColNum\{"([^"]+)"\}\s*=\s*\d+', pm)}

    # {(name, n): {"R": [r...], "G": [...], "B": [...]}}
    levels = {}
    for m in re.finditer(
            r'\$HashColData\{"([^"]+)"\}\{(\d+)\}\{([RGB])\}="([^"]*)"', pm):
        name, n, ch, vals = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        levels.setdefault((name, n), {})[ch] = [int(x) for x in vals.split(",")]

    best = {}
    for (name, n), ch in levels.items():
        if len(ch.get("R", [])) != n or len(ch.get("G", [])) != n \
                or len(ch.get("B", [])) != n:
            continue
        if name not in best or n > best[name][0]:
            best[name] = (n, ch)

    builtin = {}
    for name, (n, ch) in best.items():
        colors = ["#%02X%02X%02X" % (ch["R"][i], ch["G"][i], ch["B"][i])
                  for i in range(n)]
        builtin[name] = {
            "max": maxc.get(name, n),
            "qualitative": name in qual,
            "colors": colors,
        }
    for name, n in maxc.items():                      # keep engine-known names
        builtin.setdefault(name, {"max": n, "qualitative": name in qual,
                                  "colors": []})

    files = {}
    pdir = ROOT / "ColorsBrewer"
    if pdir.is_dir():
        for f in sorted(pdir.iterdir()):
            if not f.is_file():
                continue
            cols = []
            for line in f.read_text(encoding="utf-8", errors="replace").splitlines():
                line = line.strip()
                if re.fullmatch(r"#([0-9A-Fa-f]{6})", line):
                    try:
                        binascii.unhexlify(line[1:])
                        cols.append(line.upper())
                    except binascii.Error:
                        pass
            if cols:
                files[f.name] = {"colors": cols}
    return {"builtin": builtin, "files": files}


def extract_code_keys():
    """Keys actually read from HashConfi-like structures in the Perl sources.

    Covers both the arrow form ($h->{global}{k}) and the direct subscript form
    ($h{global}{k}), which the engine also uses (e.g. colormap_conf)."""
    keys = set()
    p2 = re.compile(r'->\{[^}]*\}\{["\']?([A-Za-z_][A-Za-z0-9_\-\.]*)["\']?\}')
    p1 = re.compile(r'\$\w+->\{["\']([A-Za-z_][A-Za-z0-9_\-\.]*)["\']\}')
    p0 = re.compile(r'\$\w+\{\s*["\']?[^}"\']*["\']?\s*\}\{\s*["\']?'
                    r'([A-Za-z_][A-Za-z0-9_\-\.]*)["\']?\s*\}')
    files = list((ROOT / "bin" / "lib").glob("*.pm"))
    files += [ROOT / "bin" / "RectChr"] + list((ROOT / "bin" / "script").glob("*.pl"))
    for f in files:
        src = f.read_text(encoding="utf-8", errors="replace")
        keys |= {m.group(1) for m in p2.finditer(src)}
        keys |= {m.group(1) for m in p1.finditer(src)}
        keys |= {m.group(1) for m in p0.finditer(src)}
    keys -= {"R", "G", "B"}
    return keys


def extract_coded_defaults():
    """Parse global => {...} and ALL => {...} blocks in init_global_defaults,
    plus ||= constants from set_default_parameters."""
    src = (ROOT / "bin" / "lib" / "ConfigDataLoader.pm").read_text(encoding="utf-8", errors="replace")
    defaults = {"global": {}, "ALL": {}}
    m = re.search(r"sub init_global_defaults \{(.*?)\n\}", src, re.S)
    if m:
        body = m.group(1)
        for scope in ("global", "ALL"):
            bm = re.search(scope + r"\s*=>\s*\{(.*?)\n\t\t\}", body, re.S)
            if not bm:
                continue
            for km in re.finditer(r"['\"]?([\w\-]+)['\"]?\s*=>\s*([^,\n]+)", bm.group(1)):
                defaults[scope][km.group(1)] = km.group(2).strip().strip("'\"")
    m = re.search(r"sub set_default_parameters \{(.*?)\n\}", src, re.S)
    if m:
        for km in re.finditer(
                r'\$HashConfi->\{ALL\}\{["\']?([\w\-]+)["\']?\}\s*\|{1,2}=\s*("([^"]*)"|[\w.+\-]+)',
                m.group(1)):
            val = km.group(2)
            defaults["ALL"][km.group(1)] = val.strip('"') if val.startswith('"') else val
    return defaults


def norm_type(t):
    t = (t or "").strip()
    if not t:
        return "text"
    if "颜色" in t and "文本" in t:
        return "color"
    if t.startswith("bool"):
        return "bool"
    if t.startswith("整数"):
        return "int"
    if t.startswith("浮点") or t == "数值":
        return "float"
    if "文件路径" in t:
        return "file"
    if t == "SVG标准":
        return "text"
    return "text"


def parse_range(t):
    m = re.search(r"\((-?\d+)\s*-\s*(-?\d+)\)", t or "")
    return (int(m.group(1)), int(m.group(2))) if m else (None, None)


# xlsx "new name" -> the key the engine actually reads (code-first corrections)
NAME_FIX = {
    "cutoff_value": "cutoff_y",
    "cutoff1_value": "cutoff1_y",
    "cutoff2_value": "cutoff2_y",
}

# ratio-type params: engine multiplies only when set (>0), so unset == 1.0 effect
RATIO_DEFAULT_ONE = {
    "track_point_size", "track_text_size", "track_geom_shape_size", "label_size",
}

# ranges validated/clamped by ConfigDataLoader::validate_and_fix_parameters
CLAMP_RANGE = {
    "padding_ratio": (0, 1),
    "stroke-opacity": (0, 1),
    "fill-opacity": (0, 1),
    "track_height": (5, None),
    "colormap_nlevels": (3, 255),
}

# SVG-attr params that the engine uses numerically (not as raw text)
TYPE_FIX = {
    "stroke-width": ("float", 0, None), "strokewidth": ("float", 0, None),
    "stroke-opacity": ("float", 0, 1), "fill-opacity": ("float", 0, 1),
    "font-size": ("float", 0, None), "text-font-size": ("float", 0, None),
    "yaxis_tick_strokewidth": ("float", 0, None),
}

# dynamic/conditional defaults verified against the engine code
# (ConfigDataLoader::set_default_parameters / check_show_column_format /
#  validate_and_fix_parameters, LocalUtils tick drawing)
DEFAULT_HINTS = {
    "track_num": (
        "默认 = File1 的数值列数（总列数-2）；存在 trackN 配置段时以段数为准",
        "Default = number of value columns in File1 (cols-2); explicit trackN sections win"),
    "show_columns": (
        "默认 = File{track编号}:4（该文件存在且有数据时），否则用 File1 第(track编号+3)列",
        "Default = File{track#}:4 if that file has data, else File1 column (track#+3)"),
    "plot_type": (
        "默认继承 trackALL（trackALL 未设时为 heatmap）",
        "Inherits trackALL (heatmap when unset)"),
    "colormap_nlevels": (
        "默认继承 trackALL（默认 8）；使用 brewer 色板时取该色板色数，范围 [3,255]",
        "Inherits trackALL (8); with a brewer palette, its color count; clamped [3,255]"),
    "upper_outlier_ratio": (
        "默认 0.95；当 plot_type 非 heatmap/highlights 且 trackALL 未修改时为 1.01",
        "0.95 by default; 1.01 when plot_type is not heatmap/highlights (trackALL unchanged)"),
    "bg_stroke_color": (
        "默认 = 本层 background_color",
        "Defaults to this track's background_color"),
    "stroke-opacity": (
        "默认 = 本层 fill-opacity（均被钳制在 [0,1]）",
        "Defaults to this track's fill-opacity (both clamped to [0,1])"),
    "yaxis_tick_unit": (
        "默认按 Y 值范围自动等分",
        "Auto-divided from the Y value range"),
    "yaxis_tick_strokewidth": (
        "默认 1",
        "Default 1"),
    "stroke-width": (
        "默认 1",
        "Default 1"),
}


# GUI metadata overrides, applied to every param after assembly.
#
# Needed because the axis-tick params come from the Perl code only (they are
# auto-added by the code-vs-xlsx cross-validation below with a hard-coded
# "track"/"text"/importance-0/name-as-label stub), and because a couple of
# xlsx descriptions are misleading. Semantics verified against the engine:
#   LocalUtils::draw_yaxis_ticks (yaxis_tick_*)  /  RectChrPlot.pm (xaxis_tick_*)
PARAM_OVERRIDES = {
    "colormap_conf": {
        "label_zh": "自定义颜色文件（全局）",
    },
    # actually drives the colorbar/legend value labels, NOT axis ticks
    "axis_tick_precision": {
        "label_zh": "图例数值精度", "label_en": "Legend value precision",
        "desc_zh": "控制色带/图例数值标签的小数位数；不影响坐标轴刻度"
                   "（Y轴刻度请用 yaxis_tick_precision）",
        "desc_en": "Decimal digits of the colorbar/legend value labels; does NOT "
                   "affect axis ticks (use yaxis_tick_precision for the Y axis)",
    },
    "yaxis_tick_precision": {
        "category": "axis", "type": "int", "importance": 1, "min": 0, "max": 10,
        "label_zh": "Y轴刻度精度", "label_en": "Y tick precision",
        "desc_zh": "Y轴刻度标签的小数位数；不设时自动：YMin/YMax 均为整数取0位，"
                   "YMax 小数位超过4位取2位，其余取1位",
        "desc_en": "Decimal digits of Y axis tick labels; auto when unset: 0 if the "
                   "Y range is integral, 2 if YMax has >4 decimals, else 1",
    },
    "yaxis_tick_num": {
        "category": "axis", "type": "int", "importance": 1, "min": 2, "default": "5",
        "label_zh": "Y轴刻度数目", "label_en": "Y tick number",
        "desc_zh": "每层Y轴的刻度条数；默认5，轨道高度小于20时取2，小于90时按高度自动",
        "desc_en": "Number of Y axis ticks per track; default 5 (2 when the track is "
                   "<20px tall, auto from height below 90px)",
    },
    "yaxis_tick_unit": {
        "category": "axis", "type": "float", "importance": 1,
        "label_zh": "Y轴刻度间隔", "label_en": "Y tick interval",
        "desc_zh": "相邻Y轴刻度之间的数值间隔；不设时按Y值范围等分",
        "desc_en": "Numeric interval between Y axis ticks; auto-divided from the Y range",
    },
    "yaxis_tick_strokewidth": {
        "category": "axis", "importance": 1, "default": "1",
        "label_zh": "Y轴刻度线宽", "label_en": "Y tick stroke width",
        "desc_zh": "Y轴刻度线的宽度，默认1",
        "desc_en": "Stroke width of the Y axis ticks, default 1",
    },
    "xaxis_tick_color": {
        "category": "axis", "type": "color", "importance": 1,
        "label_zh": "X轴刻度颜色", "label_en": "X tick color",
        "desc_zh": "X轴刻度线的颜色，默认黑色",
        "desc_en": "Color of the X axis ticks, black by default",
    },
    "xaxis_tick_direction": {
        "category": "axis", "importance": 1,
        "label_zh": "X刻度方向", "label_en": "X tick direction",
        "desc_zh": "X轴刻度线朝向：Down（默认）或 Up",
        "desc_en": "X axis tick direction: Down (default) or Up",
    },

    # the xlsx source for this one carries an unusable string; the JSON had
    # been hand-patched before, so keep the good label here instead
    "xaxis_tick_show": {
        "label_zh": "x坐标显示", "label_en": "Show X-axis ticks",
    },

    # code-only params whose zh label was left as the raw key (labels/desc only)
    "track_shift_x": {
        "label_zh": "轨道水平偏移",
        "desc_zh": "整层轨道沿X方向的平移量（默认0）",
    },
    "track_shift_y": {
        "label_zh": "轨道垂直偏移",
        "desc_zh": "整层轨道沿Y方向的平移量（默认0）",
    },
    "colormap_gradient_gap": {
        "label_zh": "渐变块间距",
        "desc_zh": "相邻渐变色块之间的间距（默认0）",
        "desc_en": "Gap between adjacent gradient blocks (default 0)",
    },
    "colormap_legend_sizeratio": {
        "label_zh": "图例大小比例",
        "desc_zh": "调整图例渐变的大小比例，设为0可隐藏图例",
    },

    # ---- v0.14.6: type corrections verified against the engine ----------
    # a font-size scale ratio, NOT a color (was rendered as a color picker)
    "label_size": {
        "type": "float", "min": 0,
        "desc_zh": "轨道名称文本的字号缩放比例（默认1.0）",
        "desc_en": "Font-size scale ratio of the track label (default 1.0)",
    },
    # a presence-based boolean per the parameter contract ("bool, 有则反");
    # note the engine currently splits the value on commas and only reverses
    # the chromosomes named there (GenomicLinkPlot.pm:1825)
    "chr_order_reverse": {
        "type": "bool", "presence_only": True,
        "label_zh": "反转染色体顺序", "label_en": "Reverse chromosome order",
        "desc_zh": "反转染色体显示顺序；设置即生效，留空=不反转"
                   "（注：当前引擎按逗号分隔的染色体名解析，写 1 只影响名为 1 的染色体）",
        "desc_en": "Reverse the chromosome display order; setting it takes effect, "
                   "empty = off (note: the engine currently parses the value as a "
                   "comma-separated chromosome list, so writing 1 only affects a "
                   "chromosome literally named 1)",
    },
    # an SVG stroke color that defaults to this track's background_color
    "bg_stroke_color": {
        "type": "color",
    },
    # numeric: multiplied by the gradient scale (engine default 4.5)
    "colormap_legend_gap": {
        "type": "float", "default": "4.5",
        "label_zh": "前后两个图列的距离",
        "desc_zh": "图例与轨道之间的间距（参与渐变刻度换算，引擎默认4.5）",
        "desc_en": "Gap between the legend and the track (engine default 4.5)",
    },

    # ---- v0.14.6: labels that read as the OPPOSITE of the engine --------
    # engine shows the legend when the value is non-zero (default 1 = shown)
    "colormap_legend_show": {
        "label_zh": "显示渐变图例", "label_en": "Show gradient legend",
        "desc_zh": "是否显示颜色渐变图例：1=显示（默认），0=隐藏；图例大小为0时也不显示",
        "desc_en": "Show the gradient legend: 1 = shown (default), 0 = hidden; "
                   "also hidden when the legend size ratio is 0",
    },
    # engine hides the background only when the value is exactly 0
    "background_show": {
        "label_zh": "显示背景颜色", "label_en": "Show background",
        "desc_zh": "是否绘制该层背景：1=显示（默认），0=不绘制背景；LinkS 建议设为0",
        "desc_en": "Draw this track's background: 1 = shown (default), 0 = hidden; "
                   "the engine suggests 0 for LinkS",
    },
    # garbled label; the value is ignored (presence sets IsNumber=0)
    "as_flag": {
        "type": "bool", "presence_only": True,
        "label_zh": "该列按文本/Flag处理", "label_en": "Treat column as flags",
        "desc_zh": "把该列数据当成文本（非数值）处理；设置即生效，留空=不生效",
        "desc_en": "Treat this column as text rather than numbers; setting it "
                   "(any value) takes effect, leaving it empty disables it",
    },

    # ---- v0.14.6: presence-only flags (value is ignored by the engine) --
    # ConfigDataLoader.pm:566-570 / GenomicLinkPlot.pm:1381: the low/high
    # colors are swapped whenever the key merely EXISTS.
    "colormap_reverse": {
        "presence_only": True,
        "label_zh": "反转渐变颜色顺序", "label_en": "Reverse gradient colors",
        "desc_zh": "反转该层渐变的颜色顺序；设置即生效（写 0 也会生效），留空=不反转",
        "desc_en": "Reverse this track's gradient colours; the engine acts on the "
                   "key merely existing (writing 0 still reverses), empty = off",
    },
    # LocalUtils.pm:2054: presence suppresses the U/D tick marks
    "hist_No_UD": {
        "presence_only": True,
        "label_zh": "hist 不画上下端刻度", "label_en": "Histogram: hide U/D ticks",
        "desc_zh": "双列直方图时不绘制上下端刻度标记（内部参数）；设置即生效，留空=不生效",
        "desc_en": "Suppress the upper/lower tick marks of a two-column histogram "
                   "(internal); setting it takes effect, empty = off",
    },
    # RectChrPlot.pm:255: presence extends the end curve outwards
    "bg_end_offset": {
        "presence_only": True,
        "label_zh": "末端曲线向外扩展", "label_en": "Extend end curves",
        "desc_zh": "染色体末端曲线从X轴0开始向外扩展；设置即生效，留空=不扩展",
        "desc_en": "Extend the chromosome end curves outward; setting it takes "
                   "effect, empty = off",
    },
    # LocalUtils.pm:786: presence flips the direction (Down/Up are ignored)
    "yaxis_tick_direction": {
        "category": "axis", "importance": 1,
        "type": "bool", "choices": None, "presence_only": True,
        "label_zh": "Y刻度方向", "label_en": "Y tick direction",
        "desc_zh": "Y轴刻度线朝反方向绘制；设置即生效（Down/Up 取值引擎不看），留空=默认方向",
        "desc_en": "Flip the Y tick direction; the engine acts on the key merely "
                   "existing (its Down/Up value is ignored), empty = default",
    },
}


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True)

    # ---- sheet2: zh descriptions keyed by param name
    ws2 = wb["cheeckAA"]
    zh_desc = {}
    for row in ws2.iter_rows(values_only=True):
        vals = [str(c).strip() if c is not None else "" for c in row]
        nz = [v for v in vals if v]
        if len(nz) >= 2 and re.match(r"^[A-Za-z][\w\-]*$", nz[0]):
            zh_desc.setdefault(nz[0], nz[1])

    # ---- sheet1: main metadata table
    ws1 = wb["参数列表"]
    rows = list(ws1.iter_rows(values_only=True))
    params, aliases = [], {}
    seen = set()
    for row in rows[2:]:
        vals = [str(c).strip() if c is not None else "" for c in row]
        name = vals[5] if len(vals) > 5 else ""
        if not name:
            continue
        name = name.split("/")[0].strip()          # 'chr_zoom_region/zoom_region'
        name = NAME_FIX.get(name, name)
        name = name.replace("colormap_ legend_gap", "colormap_legend_gap")
        if not re.match(r"^[A-Za-z][\w\-\.]*$", name):   # skip garbage/placeholder rows
            continue
        if name in INTERNAL or name in DROP or name in seen:
            continue
        seen.add(name)
        scope_cell = vals[1]
        scope = ["track"] if scope_cell == "Level" else (["global"] if scope_cell == "global" else ["global", "track"])
        cat = CAT_MAP.get(vals[3], "track" if vals[3].startswith("track") else "track")
        typ, rng = norm_type(vals[8]), parse_range(vals[8])
        if name in CHOICES:
            typ = "enum"
        if name in TYPE_FIX:
            typ, lo, hi = TYPE_FIX[name]
            rng = (lo, hi)
        imp = len(vals[7]) if set(vals[7]) <= {"*"} and vals[7] else 0
        p = {
            "name": name,
            "scope": scope,
            "category": cat,
            "old_names": [vals[2]] if vals[2] and vals[2] != name else [],
            "importance": imp,
            "type": typ,
            "choices": CHOICES.get(name),
            "min": rng[0], "max": rng[1],
            "default": None,
            "desc_zh": zh_desc.get(name) or (vals[10] if len(vals) > 10 else "") or (vals[9] if len(vals) > 9 else "") or vals[4],
            "desc_en": EN.get(name, ("", ""))[1],
            "label_zh": vals[4] or (vals[9] if len(vals) > 9 else "") or name,
            "label_en": EN.get(name, (name, ""))[0],
            "plot_types": RELEVANCE.get(name, []),
        }
        params.append(p)
        for old in p["old_names"]:
            aliases[old] = name

    # ---- extra params not present in sheet1
    extras = [
        ("File1", "files", "file", "*****", True),
        ("FileX", "files", "file", "****", False),
        ("hist_No_UD", "data", "bool", "", False),
    ]
    zh_extra = {"File1": "输入数据文件FileX路径（必需，放最前）", "FileX": "第N个输入数据文件路径",
                "hist_No_UD": "hist两列数据时不画上下端刻度标记（内部参数）"}
    for name, cat, typ, imp, required in extras:
        if name not in seen:
            params.append({
                "name": name, "scope": ["global"] if name.startswith("File") else ["track"],
                "category": cat, "old_names": [], "importance": len(imp), "type": typ,
                "choices": None, "min": None, "max": None, "default": None,
                "desc_zh": zh_desc.get(name) or zh_extra.get(name, ""),
                "desc_en": EN.get(name, ("", ""))[1],
                "label_zh": zh_extra.get(name, name),
                "label_en": EN.get(name, (name, ""))[0],
                "plot_types": [], "required": required,
            })
            aliases[f"FileN"] = "FileX"

    # ---- manual alias fixes found during code cross-validation
    aliases["colormap_legend_size"] = "colormap_legend_sizeratio"
    # legacy spellings the engine never reads (kept out of params by DROP)
    aliases["strokewidth"] = "stroke-width"
    # engine migrates the legacy global key itself
    # (ConfigDataLoader.pm:147-150: zoom_region -> chr_zoom_region)
    aliases["zoom_region"] = "chr_zoom_region"
    # remaining legacy names from bin/script/ParaOld2New.pl that are safe to map
    aliases["ShowXaxis"] = "xaxis_tick_show"
    aliases["SizeGradienRatio"] = "colormap_legend_sizeratio"
    for doc, real in (("cutoff_value", "cutoff_y"), ("cutoff1_value", "cutoff1_y"),
                      ("cutoff2_value", "cutoff2_y")):
        aliases[doc] = real

    # ---- cross-validation against the Perl code
    code_keys = extract_code_keys()
    defaults = extract_coded_defaults()
    user_keys = {p["name"] for p in params}
    for scope, dd in defaults.items():
        for k, v in dd.items():
            if k in INTERNAL or k in DROP or k in ("Bin", "RealBin"):
                continue
            for p in params:
                if p["name"] == k:
                    p["default"] = v
                    if scope == "global" and "global" not in p["scope"]:
                        p["scope"] = ["global"] + p["scope"]
                    elif scope == "ALL" and "track" not in p["scope"]:
                        p["scope"] = p["scope"] + ["track"]
            if k not in user_keys:
                params.append({
                    "name": k, "scope": ["global"] if scope == "global" else ["track"],
                    "category": "track", "old_names": [], "importance": 0, "type": "text",
                    "choices": CHOICES.get(k), "min": None, "max": None, "default": v,
                    "desc_zh": zh_desc.get(k, "代码中使用的参数"), "desc_en": EN.get(k, ("", ""))[1],
                    "label_zh": k, "label_en": EN.get(k, (k, ""))[0],
                    "plot_types": RELEVANCE.get(k, []),
                })
                user_keys.add(k)

    # scope fixes verified manually in the Perl code
    SCOPE_FIX = {
        "colormap_legend_gap": ["global"], "colormap_gradient_gap": ["track"],
        "track_shift_x": ["track"], "track_shift_y": ["track"],
        "xaxis_tick_direction": ["global"], "xaxis_tick_color": ["global"],
        "yaxis_tick_num": ["track"], "yaxis_tick_precision": ["track"],
        "yaxis_tick_unit": ["track"], "yaxis_tick_strokewidth": ["track"],
        "yaxis_tick_direction": ["track"], "track_text_overlap": ["track"],
        "colormap_legend_layout": ["track"],
    }
    for p in params:
        if p["name"] in SCOPE_FIX:
            p["scope"] = SCOPE_FIX[p["name"]]

    # ---- engine-grounded defaults: clamps, ratio=1.0 semantics, dynamic hints
    by_name = {p["name"]: p for p in params}
    for k, (lo, hi) in CLAMP_RANGE.items():
        p = by_name.get(k)
        if not p:
            continue
        if p.get("min") is None or (lo is not None and p["min"] < lo):
            p["min"] = lo
        if p.get("max") is None or (hi is not None and p["max"] > hi):
            p["max"] = hi
    for k in RATIO_DEFAULT_ONE:
        p = by_name.get(k)
        if p and not p.get("default"):
            p["default"] = "1.0"
    for k, (zh, en) in DEFAULT_HINTS.items():
        p = by_name.get(k)
        if p:
            p["default_hint_zh"] = zh
            p["default_hint_en"] = en

    in_xlsx_not_code = sorted(k for k in user_keys if k not in code_keys)
    in_code_not_xlsx = sorted(k for k in code_keys if k not in user_keys and k not in INTERNAL)

    # ---- GUI metadata overrides (labels/desc/category/type/importance)
    applied = []
    for p in params:
        ov = PARAM_OVERRIDES.get(p["name"])
        if ov:
            p.update(ov)
            applied.append(p["name"])

    # sanity: every numeric default must sit inside its [min, max] clamp
    out_of_range = []
    for p in params:
        d, mn, mx = p.get("default"), p.get("min"), p.get("max")
        if d in (None, ""):
            continue
        try:
            dv = float(d)
        except (TypeError, ValueError):
            continue
        if (mn is not None and dv < mn) or (mx is not None and dv > mx):
            out_of_range.append(p["name"])

    schema = {
        "meta": {
            "tool": "RectChr", "version": "1.50",
            "xlsx": XLSX.name, "generated_by": "gui/tools/gen_schema.py",
        },
        "categories": CATEGORIES,
        "plot_types": PLOT_TYPES,
        "palettes": extract_palettes(),
        "aliases": aliases,
        "params": params,
        "validation": {
            "in_xlsx_not_in_code": in_xlsx_not_code,
            "in_code_not_in_xlsx": in_code_not_xlsx,
            "defaults_out_of_range": sorted(out_of_range),
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(schema, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"params: {len(params)}  (xlsx-derived + extras)")
    print(f"aliases: {len(aliases)}")
    print(f"in xlsx but never read by code: {in_xlsx_not_code}")
    print(f"in code but not in xlsx:        {in_code_not_xlsx}")
    print(f"overrides applied:              {sorted(applied)}")
    print(f"defaults out of range:          {sorted(out_of_range)}")
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
